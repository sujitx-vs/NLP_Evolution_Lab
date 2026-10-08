from pathlib import Path
import json
import re
import hashlib
import unicodedata

import pandas as pd

from pypdf import PdfReader
from pptx import Presentation
from docx import Document


# ============================================================
# CONFIG
# ============================================================

DATASET_DIR = Path(r"D:\dataset")

OUTPUT_DIR = Path(r"D:\processed_corpus")

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".pptx",
    ".docx",
    ".txt",
    ".md"
}


DOCUMENTS_JSONL = OUTPUT_DIR / "documents.jsonl"

CORPUS_TXT = OUTPUT_DIR / "corpus.txt"

REPORT_CSV = OUTPUT_DIR / "extraction_report.csv"


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(text):
    """
    Light cleaning only.

    We do NOT:
    - remove stopwords
    - stem
    - lemmatize
    - remove punctuation
    - lowercase everything

    Because this corpus will be used for language modeling.
    """

    if not text:
        return ""

    # Normalize Unicode
    text = unicodedata.normalize(
        "NFKC",
        text
    )

    # Remove null bytes
    text = text.replace(
        "\x00",
        " "
    )

    # Normalize line endings
    text = text.replace(
        "\r\n",
        "\n"
    )

    text = text.replace(
        "\r",
        "\n"
    )

    # Fix words split by hyphen + line break
    #
    # Example:
    #
    # transfor-
    # mer
    #
    # becomes:
    #
    # transformer
    #
    text = re.sub(
        r"(\w)-\n(\w)",
        r"\1\2",
        text
    )

    # Replace excessive spaces/tabs
    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    # Remove spaces before punctuation
    text = re.sub(
        r"\s+([,.!?;:])",
        r"\1",
        text
    )

    # Limit excessive blank lines
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    # Remove spaces around newline
    text = re.sub(
        r" *\n *",
        "\n",
        text
    )

    return text.strip()


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf(file_path):
    reader = PdfReader(
        str(file_path)
    )

    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):
        try:
            text = page.extract_text()

            if text:
                pages.append(
                    text
                )

        except Exception as error:
            print(
                f"[PDF PAGE ERROR] "
                f"{file_path.name} "
                f"page {page_number}: {error}"
            )

    return "\n\n".join(
        pages
    )


# ============================================================
# PPTX EXTRACTION
# ============================================================

def extract_pptx(file_path):
    presentation = Presentation(
        str(file_path)
    )

    slides_text = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        current_slide = []

        for shape in slide.shapes:

            if hasattr(
                shape,
                "text"
            ):
                text = shape.text.strip()

                if text:
                    current_slide.append(
                        text
                    )

        if current_slide:
            slide_text = "\n".join(
                current_slide
            )

            slides_text.append(
                slide_text
            )

    return "\n\n".join(
        slides_text
    )


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx(file_path):
    document = Document(
        str(file_path)
    )

    contents = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            contents.append(
                text
            )

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    row_text.append(
                        text
                    )

            if row_text:
                contents.append(
                    " | ".join(
                        row_text
                    )
                )

    return "\n".join(
        contents
    )


# ============================================================
# TXT / MARKDOWN EXTRACTION
# ============================================================

def extract_text_file(file_path):
    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1"
    ]

    for encoding in encodings:

        try:
            return file_path.read_text(
                encoding=encoding
            )

        except UnicodeDecodeError:
            continue

    raise UnicodeDecodeError(
        "unknown",
        b"",
        0,
        1,
        f"Could not decode {file_path}"
    )


# ============================================================
# GENERAL EXTRACTION ROUTER
# ============================================================

def extract_file(file_path):
    extension = (
        file_path
        .suffix
        .lower()
    )

    if extension == ".pdf":

        return extract_pdf(
            file_path
        )

    elif extension == ".pptx":

        return extract_pptx(
            file_path
        )

    elif extension == ".docx":

        return extract_docx(
            file_path
        )

    elif extension in {
        ".txt",
        ".md"
    }:

        return extract_text_file(
            file_path
        )

    else:

        return ""


# ============================================================
# HASH FOR EXACT DUPLICATE DETECTION
# ============================================================

def text_hash(text):
    """
    Used to detect exact duplicate extracted documents.
    """

    normalized = re.sub(
        r"\s+",
        " ",
        text.lower()
    ).strip()

    return hashlib.sha256(
        normalized.encode(
            "utf-8"
        )
    ).hexdigest()


# ============================================================
# WORD COUNT
# ============================================================

def count_words(text):
    return len(
        re.findall(
            r"\b\w+\b",
            text
        )
    )


# ============================================================
# MAIN
# ============================================================

def prepare_corpus():

    if not DATASET_DIR.exists():

        raise FileNotFoundError(
            f"Dataset not found: "
            f"{DATASET_DIR}"
        )

    # --------------------------------------------------------
    # DISCOVER FILES
    # --------------------------------------------------------

    files = [

        file

        for file
        in DATASET_DIR.rglob("*")

        if (
            file.is_file()

            and file.suffix.lower()
            in SUPPORTED_EXTENSIONS
        )

    ]

    print(
        "\n==================================="
    )
    print(
        "CORPUS PREPARATION"
    )
    print(
        "==================================="
    )

    print(
        f"Dataset directory: "
        f"{DATASET_DIR}"
    )

    print(
        f"Supported files found: "
        f"{len(files)}"
    )


    documents = []

    report = []

    seen_hashes = {}

    successful = 0

    failed = 0

    empty = 0

    duplicates = 0


    # --------------------------------------------------------
    # PROCESS FILES
    # --------------------------------------------------------

    for index, file_path in enumerate(
        files,
        start=1
    ):

        print(
            f"[{index}/{len(files)}] "
            f"{file_path.name}"
        )

        record = {

            "file_name":
                file_path.name,

            "file_path":
                str(file_path),

            "extension":
                file_path.suffix.lower(),

            "status":
                "",

            "characters":
                0,

            "words":
                0,

            "duplicate_of":
                "",

            "error":
                ""

        }


        try:

            raw_text = extract_file(
                file_path
            )

            cleaned_text = clean_text(
                raw_text
            )


            # ------------------------------------------------
            # EMPTY DOCUMENT
            # ------------------------------------------------

            if not cleaned_text:

                record["status"] = (
                    "empty"
                )

                empty += 1

                report.append(
                    record
                )

                continue


            # ------------------------------------------------
            # DOCUMENT STATISTICS
            # ------------------------------------------------

            characters = len(
                cleaned_text
            )

            words = count_words(
                cleaned_text
            )


            record["characters"] = (
                characters
            )

            record["words"] = (
                words
            )


            # ------------------------------------------------
            # EXACT DUPLICATE DETECTION
            # ------------------------------------------------

            doc_hash = text_hash(
                cleaned_text
            )


            if doc_hash in seen_hashes:

                original_file = (
                    seen_hashes[
                        doc_hash
                    ]
                )


                record["status"] = (
                    "duplicate"
                )

                record["duplicate_of"] = (
                    original_file
                )

                duplicates += 1

                report.append(
                    record
                )

                print(
                    f"    DUPLICATE OF: "
                    f"{original_file}"
                )

                continue


            seen_hashes[
                doc_hash
            ] = file_path.name


            # ------------------------------------------------
            # SAVE DOCUMENT
            # ------------------------------------------------

            document_record = {

                "id":
                    len(documents),

                "source":
                    file_path.name,

                "path":
                    str(file_path),

                "type":
                    file_path
                    .suffix
                    .lower()
                    .replace(
                        ".",
                        ""
                    ),

                "characters":
                    characters,

                "words":
                    words,

                "text":
                    cleaned_text

            }


            documents.append(
                document_record
            )


            record["status"] = (
                "success"
            )

            successful += 1


        except Exception as error:

            record["status"] = (
                "failed"
            )

            record["error"] = str(
                error
            )

            failed += 1

            print(
                f"    ERROR: {error}"
            )


        report.append(
            record
        )


    # ========================================================
    # SAVE JSONL
    # ========================================================

    with open(
        DOCUMENTS_JSONL,
        "w",
        encoding="utf-8"
    ) as file:

        for document in documents:

            file.write(

                json.dumps(
                    document,
                    ensure_ascii=False
                )

                + "\n"
            )


    # ========================================================
    # SAVE COMBINED CORPUS
    # ========================================================

    with open(
        CORPUS_TXT,
        "w",
        encoding="utf-8"
    ) as file:

        for document in documents:

            file.write(
                "\n\n"
            )

            file.write(
                "<DOCUMENT_START>\n"
            )

            file.write(
                document["text"]
            )

            file.write(
                "\n<DOCUMENT_END>"
            )

            file.write(
                "\n\n"
            )


    # ========================================================
    # SAVE REPORT
    # ========================================================

    report_df = pd.DataFrame(
        report
    )


    report_df.to_csv(

        REPORT_CSV,

        index=False,

        encoding="utf-8"
    )


    # ========================================================
    # SUMMARY
    # ========================================================

    total_characters = sum(

        document["characters"]

        for document
        in documents

    )


    total_words = sum(

        document["words"]

        for document
        in documents

    )


    print(
        "\n==================================="
    )
    print(
        "CORPUS PREPARATION COMPLETE"
    )
    print(
        "==================================="
    )

    print(
        f"Successfully extracted: "
        f"{successful}"
    )

    print(
        f"Empty documents: "
        f"{empty}"
    )

    print(
        f"Exact duplicates removed: "
        f"{duplicates}"
    )

    print(
        f"Failed files: "
        f"{failed}"
    )

    print(
        f"Final documents: "
        f"{len(documents)}"
    )

    print(
        f"Total characters: "
        f"{total_characters:,}"
    )

    print(
        f"Total words: "
        f"{total_words:,}"
    )

    print(
        "\nOutput files:"
    )

    print(
        f"Documents JSONL: "
        f"{DOCUMENTS_JSONL}"
    )

    print(
        f"Combined corpus: "
        f"{CORPUS_TXT}"
    )

    print(
        f"Extraction report: "
        f"{REPORT_CSV}"
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    prepare_corpus()